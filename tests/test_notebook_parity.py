"""Run original notebook cells as an independent conversion/export oracle."""

import argparse
import contextlib
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch

import cfgrib
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
import pandas as pd
import xarray as xr

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from grib_to_csv import convert_grib, dataset_to_dataframe

SAMPLE = None
ARTIFACTS = ROOT / "data" / "parity"
NOTEBOOK = ROOT / "GribToCsv.ipynb"
REPORT = []

# Full header properties observed in the supplied reference workbook. These
# constants keep regressions portable; the user's Downloads files are optional.
REFERENCE_HEADER_STYLE = {
    "font": Font(bold=True),
    "alignment": Alignment(horizontal="center", vertical="top"),
    "border": Border(left=Side(style="thin"), right=Side(style="thin"),
                     top=Side(style="thin"), bottom=Side(style="thin")),
    "fill": PatternFill(),
    "number_format": "General",
    "protection": Protection(locked=True, hidden=False),
}


def notebook_reference(input_path, directory, monthly, dataset=None, filters=None):
    """Execute helper, original import/filter, processing, and export cells.

    Only local path variables and disabling cfgrib's index sidecar are adapted.
    Synthetic datasets replace file opening; processing sources stay untouched.
    """
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    directory.mkdir(parents=True, exist_ok=True)
    namespace = {"input_filePath": str(input_path), "raw_file_name": input_path.stem,
                 "input_directory": str(directory)}
    original_open = cfgrib.open_dataset

    def open_reference(*args, **kwargs):
        if dataset is not None:
            return dataset.copy(deep=True)
        kwargs["indexpath"] = ""
        return original_open(*args, **kwargs)

    # Cell 7 has hardcoded filter options; inject only the selected input filters.
    import_cell = 9
    if filters:
        namespace.update(prefix=filters["dataType"],
                         prefix_numberOfPoints=filters.get("numberOfPoints", -1))
        import_cell = 7
    route = [23, 24, 25] if monthly else [11, 12, 13, 14]
    with contextlib.redirect_stdout(io.StringIO()), patch.object(cfgrib, "open_dataset", open_reference):
        for index in [3, import_cell, *route]:
            source = "".join(notebook["cells"][index]["source"])
            if index == 7:
                # Input selection only; no processing or export code is changed.
                source = source.replace("prefix = 'es'", f"prefix = {filters['dataType']!r}")
                source = source.replace("prefix_numberOfPoints = -1",
                                        f"prefix_numberOfPoints = {filters.get('numberOfPoints', -1)}")
            exec(compile(source, f"{NOTEBOOK.name}:cell{index}", "exec"), namespace)
    namespace["dataset"].close()
    frame = namespace["df_complete" if monthly else "complete_df"]
    stem = Path(namespace["output_filePath"])
    return frame, Path(f"{stem}.csv"), Path(f"{stem}.xlsx")


def worksheet_contents(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        return [(sheet.title, sheet.max_row, sheet.max_column,
                 [[(cell.value, cell.data_type) for cell in row]
                  for row in sheet.iter_rows()]) for sheet in workbook]
    finally:
        workbook.close()


class NotebookParity(unittest.TestCase):
    def assert_parity(self, label, reference, actual):
        expected_frame, expected_csv, expected_xlsx = reference
        actual_frame, actual_csv, actual_xlsx = actual
        pd.testing.assert_frame_equal(expected_frame, actual_frame, check_exact=True)
        actual_bytes = actual_csv.read_bytes()
        # Original notebook keeps the platform's newlines; only normalize CRLF.
        self.assertEqual(expected_csv.read_bytes().replace(b"\r\n", b"\n"), actual_bytes)
        self.assertNotIn(b"\r", actual_bytes)
        self.assertTrue(actual_bytes.endswith(b"\n"))
        self.assertEqual(actual_bytes.count(b"\n"), len(actual_frame) + 1)
        self.assertEqual(worksheet_contents(expected_xlsx), worksheet_contents(actual_xlsx))
        workbook = openpyxl.load_workbook(actual_xlsx)
        try:
            for cell in workbook.active[1]:
                for property_name, expected in REFERENCE_HEADER_STYLE.items():
                    with self.subTest(case=label, cell=cell.coordinate, property=property_name):
                        self.assertEqual(getattr(cell, property_name), expected)
        finally:
            workbook.close()
        REPORT.append({"case": label, "rows": len(actual_frame),
                       "columns": list(actual_frame.columns),
                       "dtypes": [str(dtype) for dtype in actual_frame.dtypes],
                       "csv_sha256": hashlib.sha256(actual_csv.read_bytes()).hexdigest(),
                       "dataframe_exact": True, "csv_bytes_exact_after_newline_normalization": True,
                       "csv_lf_only": True, "xlsx_worksheet_exact": True,
                       "xlsx_reference_header_style_exact": True})

    def test_sample_yearly_and_cli(self):
        if SAMPLE is None:
            self.skipTest("Supply --sample for real GRIB and CLI parity")
        directory = ARTIFACTS / "sample_yearly"
        reference = notebook_reference(SAMPLE, directory / "notebook", monthly=False)
        actual = convert_grib(SAMPLE, directory / "script")
        self.assert_parity("sample_yearly", reference, actual)
        cli_dir = directory / "cli"
        subprocess.run([sys.executable, str(ROOT / "grib_to_csv.py"), str(SAMPLE),
                        "--output-dir", str(cli_dir)], check=True, capture_output=True, text=True)
        self.assertEqual(actual[1].read_bytes(), (cli_dir / actual[1].name).read_bytes())
        self.assertEqual(worksheet_contents(actual[2]), worksheet_contents(cli_dir / actual[2].name))

    def test_sample_filters(self):
        if SAMPLE is None:
            self.skipTest("Supply --sample for real GRIB filter parity")
        import eccodes
        with SAMPLE.open("rb") as handle:
            message = eccodes.codes_grib_new_from_file(handle)
            try:
                filters = {"dataType": eccodes.codes_get(message, "dataType"),
                           "numberOfPoints": eccodes.codes_get(message, "numberOfPoints")}
            finally:
                eccodes.codes_release(message)
        for selected in ({"dataType": filters["dataType"]}, filters):
            label = "sample_filter_" + "_".join(str(v) for v in selected.values())
            directory = ARTIFACTS / label
            reference = notebook_reference(SAMPLE, directory / "notebook", False, filters=selected)
            actual = convert_grib(SAMPLE, directory / "script",
                                  data_type=selected["dataType"],
                                  number_of_points=selected.get("numberOfPoints"))
            self.assert_parity(label, reference, actual)

    def test_synthetic_routes_with_missing_values_and_dtype_promotion(self):
        for monthly in (False, True):
            coordinates = {"time": np.array(["2024-12-31T21:00", "2025-01-01T03:00"],
                                             dtype="datetime64[ns]"),
                           "latitude": [5.0, -7.5], "longitude": [110.0, 111.0]}
            dimensions = ("time", "latitude", "longitude")
            shape = (2, 2, 2)
            if monthly:
                coordinates["step"] = np.array([0, 90, 2940], dtype="timedelta64[m]")
                dimensions = ("time", "step", "latitude", "longitude")
                shape = (2, 3, 2, 2)
            first = np.arange(np.prod(shape), dtype=np.float32).reshape(shape)
            first.flat[-1] = np.nan
            second = (first.astype(np.float64) + 0.123456789)
            dataset = xr.Dataset({"first": (dimensions, first), "second": (dimensions, second)},
                                 coords=coordinates)
            label = "synthetic_monthly" if monthly else "synthetic_yearly"
            directory = ARTIFACTS / label
            placeholder = directory / "synthetic.grib"
            directory.mkdir(parents=True, exist_ok=True)
            placeholder.touch()
            reference = notebook_reference(placeholder, directory / "notebook", monthly, dataset)
            with patch.object(cfgrib, "open_dataset", return_value=dataset.copy(deep=True)):
                actual = convert_grib(placeholder, directory / "script")
            self.assert_parity(label, reference, actual)
            if monthly:
                self.assertEqual(actual[0]["hour"].tolist(), [0, 1, 49, 0, 1, 49])
                self.assertEqual(actual[0]["day"].tolist(), [31, 31, 31, 1, 1, 1])
            self.assertTrue(np.isnan(actual[0]["first"].iloc[-1]))

    def test_rejects_unsupported_dimensions_and_empty_filters(self):
        scalar = xr.Dataset({"temperature": xr.DataArray(1.0)})
        with self.assertRaisesRegex(ValueError, "Unsupported GRIB dimensions"):
            dataset_to_dataframe(scalar)
        with self.assertRaisesRegex(ValueError, "No GRIB variables"):
            dataset_to_dataframe(xr.Dataset())


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sample", type=Path, help="Local yearly GRIB sample")
    options = parser.parse_args()
    SAMPLE = options.sample.resolve() if options.sample else None
    original_hash = hashlib.sha256(NOTEBOOK.read_bytes()).hexdigest()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(NotebookParity)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if hashlib.sha256(NOTEBOOK.read_bytes()).hexdigest() != original_hash:
        raise AssertionError("Original notebook was modified")
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    report_path = ARTIFACTS / "parity_report.json"
    report_path.write_text(json.dumps({"notebook_sha256": original_hash, "cases": REPORT,
                                      "passed": result.wasSuccessful()}, indent=2), encoding="utf-8")
    print(f"Parity report: {report_path}")
    raise SystemExit(0 if result.wasSuccessful() else 1)
